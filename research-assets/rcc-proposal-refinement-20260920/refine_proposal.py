from pathlib import Path
from copy import deepcopy
from hashlib import sha256
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

SOURCE = Path(r'[local path redacted]')
OUT = Path(r'[local path redacted]')
ROOT = Path(__file__).parent
original_hash = sha256(SOURCE.read_bytes()).hexdigest()
doc = Document(SOURCE)
original = list(doc.paragraphs)

texts = {
0: 'MatterSyn for Predicting Colloidal Nanocrystal Synthesis',
1: 'Research goals and significance',
2: 'Quantum dots already enable commercial display technologies [1], while nanodiamonds with optically active defects are being developed for quantum sensing and computing [2]. These applications depend on controlling particle size, morphology, surfaces, and defects. Achieving that control remains difficult, including in strongly covalent materials such as diamond. MatterSyn addresses the synthesis problem: how can we turn a desired material specification into a practical recipe? Our initial focus is colloidal semiconductor nanocrystals. The broader motivation is to make useful nanomaterials more accessible for semiconductor, optoelectronic, and quantum technologies.',
3: 'We will test whether learning from structured, source-linked synthesis records produces more useful recipes than retrieving similar papers or generating instructions directly from their text. The model will propose precursors and processing conditions for a specified composition, crystal phase, and particle size, while identifying missing information and uncertain predictions. We will build the dataset, compare compact models, and test selected proposals experimentally in Professor Paul Alivisatos\'s laboratory. The research will deliver a curated dataset, reproducible model comparisons, and experimentally evaluated recipes. The public MatterSyn website will let readers inspect the evidence behind each contribution.',
4: 'Existing foundation and planned scope',
5: 'The public MatterSyn pilot [3] contains 675 structured records from 41 source contributions, including 123 synthesis routes or variants across 39 direct synthesis systems. The remaining records include supporting procedures and characterization observations, so this total is not a count of independent experiments. The pilot establishes a working data format and a way to link recipes to source evidence. Our longer-term target is approximately 13,500 structured records, about twenty times the current collection; the number suitable for model training will depend on their completeness and review.',
6: 'Previous RCC use',
7: 'This is the project\'s first RCC allocation request. The existing MatterSyn pilot was developed outside RCC, and there are no results or publications from previous RCC allocations to report. We have not yet measured performance on Midway.',
9: 'Computational research plan',
10: 'Dataset construction. We will screen papers for usable synthesis information, remove duplicates, and match each main article to its supporting information using document content. Detailed extraction will preserve chemical identities, amounts, concentrations, operation order, temperature, atmosphere, workup, and sample-specific characterization. Language and vision models will assist with extraction; a separate audit will check each record against the source. Ambiguous assignments will receive human review, and missing quantities will remain missing. A manually checked benchmark of 200 paper groups will measure extraction accuracy, including numerical values, units, operation order, and sample assignment. Unaudited records will stay outside the training and public datasets; unreadable or unmatched sources will remain unresolved.',
11: 'Recipe prediction and evaluation. We will compare literature retrieval, generation from retrieved text, and parameter-efficient adaptation of open-weight models with 1–8 billion parameters. The models will return candidate recipes, sources, and uncertainty estimates. All comparisons will use a common training and retrieval corpus. Evaluation papers, their supporting information, and closely related recipe variants will be withheld from both training and retrieval. We will assess agreement with reported precursors and conditions, chemical plausibility, source support, and whether instructions are complete enough to execute. Agreement with a published recipe is a retrospective measure; experiments will test whether proposed recipes achieve their targets. The current dataset has no verified pairs linking measured atomic coordinates to complete recipes; exact structure-conditioned training will require such pairs.',
12: 'Experimental feedback. We plan to test predictions in Professor Paul Alivisatos\'s laboratory alongside model development. Targets and success criteria, such as phase, particle size, or optical response, will be selected before inspecting model outputs. We will compare model proposals with literature-based starting recipes, record unsuccessful outcomes, and repeat promising conditions to assess reproducibility. These results will inform subsequent model updates. A separate set of targets will evaluate the updated model, so experiments used for fine-tuning do not also serve as its final test.',
13: 'Efficient execution and milestones',
14: 'Months 1–2 will focus on screening and dataset construction. During the first two weeks, representative short, long, scanned, and supporting-information-heavy papers will establish runtimes, memory use, and charged SUs. Papers will run independently through Slurm job arrays. We will reuse extracted text, batch inference by input length, and checkpoint completed work. The first release will prioritize sources with complete synthesis and sample-resolved characterization; its size will depend on audit capacity.',
15: 'Months 3–6 will focus on model comparisons and prospective laboratory tests. The training budget allows 20 adaptation and evaluation runs covering model size, data representation, and repeated seeds. Most jobs will use one GPU; two-GPU jobs will be used only if benchmarks justify them.',
16: 'Months 7–12 will incorporate experimental feedback, assess the revised models on reserved targets, and release the reviewed data and results. We will track GPU utilization and SU consumption throughout, adjusting concurrency and batch size to fit the allocation.',
18: 'Compute request and justification',
19: 'We request 375,000 SUs for the staged workload below. Screening will use metadata and available text from up to 50,000 paper groups. Detailed parsing and model extraction will cover up to 10,000 selected groups, with independent audits budgeted for up to 4,000. These are capacity limits, not forecasts of how many papers contain usable recipes. GPU resources will support repeated language and vision model inference and adapter training; CPU resources will handle document processing and evaluation. All runtimes are planning estimates to be tested in the initial benchmark.',
20: 'SU estimate. Using provisional allowances of 1 SU per CPU core-hour and 100 SUs per GPU-hour, including the associated host resources, the base estimate is 26,000 + 300,000 = 326,000 SUs. Adding 15% for reruns and variation between papers gives 374,900 SUs, rounded to a request of 375,000. These allowances are not verified Midway3 billing rates. RCC charges depend on the resources and partition used [4]; we will confirm the conversion with RCC before submission. No measured throughput or scaling result is claimed.',
21: 'Job configuration. CPU jobs will request 4–8 cores and 16–32 GB RAM on one node. GPU inference jobs will request one GPU, 4–8 host cores, and 32–64 GB host RAM. Adapter training may use two GPUs on one node and 64–128 GB host RAM. We target 24–48 GB of GPU memory, using smaller or quantized models when necessary. Initial concurrency will be limited to four CPU jobs and four single-GPU jobs, subject to availability. We do not require multi-node training or 768 GB big-memory nodes. GPU-hours count every allocated device.',
22: 'Storage and research outputs',
23: 'We request 1,000 GB of persistent storage: 150 GB for working copies of selected sources and metadata, 150 GB for structured records and selected evidence, 250 GB for model weights and software environments, 350 GB for checkpoints and evaluation outputs, and 100 GB reserve. The full source archive will remain outside RCC. Temporary page images and intermediate files will use up to approximately 1 TB of scratch space, subject to RCC policy. We will process documents in batches and remove intermediates after validation.',
24: 'Reviewed data, code, citations, and model artifacts will be released where licensing permits. Original papers and supporting information will remain access-controlled. Each released recipe will retain its source and review history, allowing readers to check the result and researchers to reproduce the model comparisons.',
25: 'References',
}

leadins = ['Dataset construction.', 'Recipe prediction and evaluation.', 'Experimental feedback.', 'SU estimate.', 'Job configuration.']
for index, text in texts.items():
    p = original[index]
    p.clear()
    lead = next((s for s in leadins if text.startswith(s)), None)
    if lead:
        p.add_run(lead).bold = True
        p.add_run(text[len(lead):])
    else:
        p.add_run(text)

scope = original[6].insert_paragraph_before('This request supports screening up to 50,000 candidate paper and supporting-information groups, detailed extraction from up to 10,000, and independent audit of up to 4,000. Screening will identify the most useful sources; it will not be reported as complete review. We will publish and train only on audited records. The allocation funds this defined stage rather than exhaustive curation of the full collection.')

# Preserve the uploaded document's structure and page geometry, while repairing its text styles.
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Header', 'Footer']:
    st=doc.styles[name]
    st.font.name='Calibri'
    st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.04
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.widow_control=True
doc.styles['Title'].font.size=Pt(20)
doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 1'].font.size=Pt(12)
doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
doc.styles['Heading 1'].paragraph_format.keep_with_next=True
for st in doc.styles:
    for e in list(st.element.iter(qn('w:pBdr'))): e.getparent().remove(e)
for p in doc.paragraphs:
    for e in list(p._p.iter(qn('w:pBdr'))): e.getparent().remove(e)

table=doc.tables[0]
newrow=table.add_row()
table.rows[0]._tr.addnext(newrow._tr)
rows=[
 ['Workload','Planning calculation','Raw resource demand'],
 ['Initial screening','50,000 groups × 0.01 h × 4 cores','2,000 CPU core h'],
 ['Selected parsing and OCR','10,000 groups × 0.5 h × 4 cores','20,000 CPU core h'],
 ['Indexing and evaluation','500 jobs × 1 h × 8 cores','4,000 CPU core h'],
 ['Model extraction','10,000 groups × 0.15 GPU h','1,500 GPU h'],
 ['Independent source audit','4,000 groups × 0.25 GPU h','1,000 GPU h'],
 ['Adaptation and evaluation','20 runs × 25 GPU h','500 GPU h'],
 ['Total','CPU and GPU workloads counted separately','26,000 CPU core h\n3,000 GPU h'],
]
widths=[2.05,2.88,2.07]
table.autofit=False
for c,w in zip(table.columns,widths): c.width=Inches(w)
pr=table._tbl.tblPr
old=pr.find(qn('w:tblBorders'))
if old is not None: pr.remove(old)
borders=OxmlElement('w:tblBorders')
for side in ['top','bottom','left','right','insideH','insideV']:
    e=OxmlElement('w:'+side)
    for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
    borders.append(e)
pr.append(borders)
for i,(row,vals) in enumerate(zip(table.rows,rows)):
    trpr=row._tr.get_or_add_trPr()
    if trpr.find(qn('w:cantSplit')) is None: trpr.append(OxmlElement('w:cantSplit'))
    if i==0 and trpr.find(qn('w:tblHeader')) is None: trpr.append(OxmlElement('w:tblHeader'))
    for j,(cell,value) in enumerate(zip(row.cells,vals)):
        cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text=value
        tcpr=cell._tc.get_or_add_tcPr()
        for name in ['tcMar','shd']:
            e=tcpr.find(qn('w:'+name))
            if e is not None: tcpr.remove(e)
        margins=OxmlElement('w:tcMar')
        for side,n in [('top',70),('bottom',70),('left',85),('right',85)]:
            e=OxmlElement('w:'+side);e.set(qn('w:w'),str(n));e.set(qn('w:type'),'dxa');margins.append(e)
        tcpr.append(margins)
        if i==0:
            e=OxmlElement('w:shd');e.set(qn('w:fill'),'E8E8E8');tcpr.append(e)
        p=cell.paragraphs[0]
        p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j==2 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:r.font.size=Pt(10);r.bold=i in [0,7]
original[20].paragraph_format.space_before=Pt(7)

def link(p,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    for tag,val in [('color','000000'),('u','single'),('sz','18')]:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);pr.append(e)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

refs=[
 ('[1] Nanosys. ','Quantum Dots Take Center Stage at Display Week 2025','https://www.nanosys.com/blog-newsroom/quantum-dots-take-center-stage-at-display-week-2025','. 30 May 2025.'),
 ('[2] Liang J et al. ','Bottom-up synthesis of molecular nanodiamond from nanographene','https://doi.org/10.1038/s41586-026-10669-3','. Nature 655, 102–108 (2026).'),
 ('[3] MatterSyn. ','Materials synthesis atlas and dataset','https://cuiyist.github.io/mattersyn-site/','. Version 0.33.0, 20 September 2026.'),
 ('[4] University of Chicago RCC. ','Running jobs and service units','https://docs.rcc.uchicago.edu/slurm/main/','; '),
]
for i,(lead,label,url,tail) in enumerate(refs):
    p=original[26+i] if i<3 else doc.add_paragraph()
    p.clear();p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.0
    p.add_run(lead).font.size=Pt(9);link(p,label,url);p.add_run(tail).font.size=Pt(9)
    if i==3:link(p,'allocation requirements','https://rcc.uchicago.edu/accounts-allocations/research-ii-allocation-request')

# Review comments flag unresolved assumptions without interrupting the submission text.
doc.add_comment(scope.runs, 'Scope reconciled to your choice: screen up to 50,000 groups, extract up to 10,000, and audit up to 4,000. The added screening row costs an assumed 2,000 CPU core-hours. Only audited records enter training and public releases.', author='Editorial review', initials='ER')
doc.add_comment(original[5].runs, 'The twenty-fold expansion is retained as a longer-term target of about 13,500 structured records, not a measured yield or a count of independent experiments. A pilot estimate of eligible-paper yield would strengthen this target.', author='Editorial review', initials='ER')
doc.add_comment(original[12].runs, 'Before submission, specify a feasible initial number of prospective targets, replicate batches, and the characterization methods available in the Alivisatos lab. These details would make the validation plan more concrete; no experimental capacity has been invented here.', author='Editorial review', initials='ER')
doc.add_comment(original[20].runs, 'Most important remaining technical check: confirm the intended Midway partition and effective CPU/GPU SU rates with RCC. The 1 and 100 SU conversion factors and all runtimes are assumptions, not measured tariffs or benchmarks. A representative pilot should report model size, input length, GPU type/memory, elapsed time, and charged SUs.', author='Editorial review', initials='ER')

# Ensure edits do not retain imported low control characters in the body or references.
for p in doc.paragraphs:
    assert not any(0x80 <= ord(c) <= 0x9f or ord(c)==65533 for c in p.text), repr(p.text)
doc.core_properties.title=texts[0]
doc.core_properties.subject='Refined RCC Research II proposal for staged synthesis data curation and model validation'
doc.save(OUT)
assert sha256(SOURCE.read_bytes()).hexdigest()==original_hash
(ROOT/'source-sha256.txt').write_text(original_hash+'  '+str(SOURCE)+'\n',encoding='utf-8')
print('Saved',OUT)
print('Source unchanged:',original_hash)
print('Bytes:',OUT.stat().st_size)
print('Words:',sum(len(p.text.split()) for p in doc.paragraphs)+sum(len(c.text.split()) for t in doc.tables for row in t.rows for c in row.cells))
print('Review comments:',len(doc.comments))
