from pathlib import Path
from shutil import copy2
from docx import Document

root=Path(__file__).parent
path=Path(r'[local path redacted]')
backup=root/'MatterSyn_RCC_Research_II_Proposal_Refined_before_academic_intro.docx'
if not backup.exists():copy2(path,backup)
d=Document(path)
old=list(d.paragraphs)
intro='Nanocrystals have broad technological applications, ranging from semiconductor quantum dots in commercial displays [1] to nanodiamonds under investigation for quantum sensing and computing [2]. However, the reproducible synthesis of nanocrystals with prescribed size, morphology, crystal phase, and defect characteristics remains challenging, particularly for strongly covalent materials such as diamond. Predictive synthesis methods could expand the range of accessible nanomaterials and support advances in semiconductor, optoelectronic, and quantum technologies.'
definition='To address this challenge, we are developing MatterSyn, a materials synthesis data platform that combines a curated, machine-readable database of experimental procedures and characterization results with an interactive materials atlas [3]. Reported precursor chemistry, processing conditions, and product characteristics are linked to the corresponding papers and available supporting information. This connection between synthesis, structure, and properties provides traceable data for both experimental researchers and computational model development.'
aims='The proposed research will use MatterSyn to train and validate compact neural network models that recommend precursors and processing conditions for specified material compositions, crystal phases, and particle sizes. We will test whether structured synthesis records improve recipe prediction relative to literature retrieval and generation from unstructured text. Our initial focus is colloidal semiconductor nanocrystals, with selected predictions evaluated experimentally in Professor Paul Alivisatos\'s laboratory. The expected outcomes are a curated dataset, reproducible model comparisons, and experimentally evaluated synthesis recipes.'
assert old[2].text.startswith('Quantum dots already enable')
assert old[3].text.startswith('We will test whether learning')
old[2].clear();old[2].add_run(intro)
old[3].clear();old[3].add_run(definition)
old[4].insert_paragraph_before(aims)
d.save(path)

# Keep the reproducible refinement script synchronized with the final prose.
builder=root/'refine_proposal.py'
source=builder.read_text(encoding='utf-8')
lines=source.splitlines()
for idx,text in [(2,intro),(3,definition)]:
    matches=[i for i,line in enumerate(lines) if line.startswith(str(idx)+': ')]
    assert len(matches)==1
    lines[matches[0]]=str(idx)+': '+repr(text)+','
source='\n'.join(lines)+'\n'
anchor='scope = original[6].insert_paragraph_before('
assert source.count(anchor)==1
source=source.replace(anchor,'original[4].insert_paragraph_before('+repr(aims)+')\n\n'+anchor)
builder.write_text(source,encoding='utf-8')

new=Document(path)
assert len(new.paragraphs)==len(old)+1
assert [p.text for p in old[4:]]==[p.text for p in new.paragraphs[5:]]
assert len(new.comments)==4
print('Updated introduction with MatterSyn definition; all later paragraphs unchanged.')
print('Bytes',path.stat().st_size)
