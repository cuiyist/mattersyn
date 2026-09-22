from pathlib import Path
from hashlib import sha256
from docx import Document

ROOT = Path(__file__).parent
SOURCE = Path(r'[local path redacted]')
ORIGINAL = Path(r'[local path redacted]')
OUT = Path(r'[local path redacted]')
source_hash = sha256(SOURCE.read_bytes()).hexdigest()
original_hash = sha256(ORIGINAL.read_bytes()).hexdigest()
doc = Document(SOURCE)

replacements = {
    'Dataset construction.': 'Dataset construction. We will remove duplicates, match each article to its supporting information, and identify usable synthesis procedures. Language and vision models will extract starting chemicals, quantities, operations, reaction conditions, purification steps, and measurements of the resulting particles. A separate audit will check records against the sources, with human review of ambiguous cases. A manually checked set of 200 paper groups will be divided in advance into development and held-out evaluation subsets. Checks will cover numerical values, units, and links between recipes and measured samples. Missing information will remain explicit; only audited records will enter training and public release.',
    'Months 1–2 will': 'During months 1–2, we will construct the dataset and measure processing costs. The first two weeks will benchmark papers of different lengths and formats using the development subset. We will compare model sizes and batch sizes, recording elapsed time, peak CPU and GPU memory, utilization, and charged SUs. The decision criterion will be cost per accepted, audited record at a common accuracy requirement, including retries and host resources. The evaluation subset will remain untouched during tuning. Benchmark costs are included in the workload below; measured costs will determine the affordable production scope.',
    'Months 3–6 will': 'CPU screening will precede GPU extraction. We will cache outputs with source and processing-version identifiers, apply OCR where text extraction fails, and use vision models where text alone cannot recover the evidence. Independent papers will run in Slurm job arrays, with GPU jobs starting only after their CPU inputs are ready. Batching similar input lengths will reduce padding and repeated model loading. Completed stages will be checkpointed so failures do not require restarting an entire paper. New extraction will be limited by audit capacity, avoiding a growing backlog of unusable training records.',
    'Months 7–12 will': 'Months 3–6 will compare recipe models and begin laboratory tests; months 7–12 will incorporate experimental feedback and evaluate reserved targets. The budget supports up to 20 adaptation and evaluation runs. Validation-based early stopping will discontinue weak configurations, reserving repeated seeds for promising comparisons. Monthly reviews of charged SUs and accepted records will guide further work. If costs exceed estimates, we will reduce lower-priority extraction or model comparisons while preserving source audits and final evaluation.',
    'We request 375,000 SUs': 'We request 375,000 SUs for the staged workload below. CPU resources support source screening, document preparation, indexing, and evaluation; GPUs support extraction, independent source checking, and model adaptation. These stages produce the reviewed training data and model comparisons needed to test our research hypothesis. Workload sizes are upper limits, and runtimes are planning estimates to be tested in the initial benchmark.',
    'Job configuration.': 'Job configuration. Initial estimates are 4–8 cores and 16–32 GB RAM for CPU jobs, and one GPU with 4–8 host cores and 32–64 GB host RAM for inference. Training may need 64–128 GB host RAM. We target 24–48 GB GPU memory and will adjust requests from measured peaks with headroom. Single-node, single-GPU execution is the default. Two GPUs will be used only when needed for model memory or when benchmarks show lower charged SUs for equivalent work and accuracy. No multi-node training or 768 GB nodes are required. GPU-hours count all allocated devices.',
    'We request 1,000 GB': 'We request 1,000 GB of persistent storage, provisionally divided into 150 GB for selected sources and metadata, 150 GB for records and evidence, 250 GB for shared model weights and environments, 350 GB for checkpoints and evaluations, and 100 GB reserve. The full source archive remains outside RCC. Temporary images and intermediates will use up to approximately 1 TB of scratch, subject to RCC policy. Batches will bound scratch use; verified intermediates will be removed when reproducible from retained sources. Shared base models, compact adapters, and retention of the best and latest resumable checkpoints will limit duplication.',
}

def replace_paragraph(p, text):
    p.clear()
    lead = next((s for s in ('Dataset construction.', 'Job configuration.') if text.startswith(s)), None)
    if lead:
        p.add_run(lead).bold = True
        p.add_run(text[len(lead):])
    else:
        p.add_run(text)

for prefix, text in replacements.items():
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(matches) == 1, (prefix, len(matches))
    replace_paragraph(matches[0], text)

benchmark_paragraph = next(p for p in doc.paragraphs if p.text.startswith('During months 1–2'))
doc.add_comment(benchmark_paragraph.runs, 'Efficiency is presented as an execution and measurement plan, not as an observed Midway result. Before submission, any available pilot measurements would strengthen this section. Set field-level accuracy requirements using development data before comparing resource configurations; preserve the held-out evaluation subset. Include unsuccessful jobs and associated host resources when measuring SUs per accepted record.', author='Editorial review', initials='ER')

doc.core_properties.subject = 'RCC Research II proposal with staged resource use and quality-controlled efficiency measurements'
doc.save(OUT)
assert sha256(SOURCE.read_bytes()).hexdigest() == source_hash
assert sha256(ORIGINAL.read_bytes()).hexdigest() == original_hash
assert original_hash == '88fff2efc620561e862f2580dfe7d682121e6112da73386d423b1b7a7413c1c4'
assert 26000 + 3000 * 100 == 326000
assert 326000 * 1.15 == 374900
assert 150 + 150 + 250 + 350 + 100 == 1000
print('Saved:', OUT)
print('Original and prior revision unchanged')
print('Bytes:', OUT.stat().st_size)
print('Comments:', len(doc.comments))
print('Words:', sum(len(p.text.split()) for p in doc.paragraphs) + sum(len(c.text.split()) for t in doc.tables for r in t.rows for c in r.cells))
