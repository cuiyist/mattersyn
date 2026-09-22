from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(r'[local path redacted]')
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.64)
sec.bottom_margin = Inches(.62)
sec.left_margin = sec.right_margin = Inches(.75)
sec.header_distance = Inches(.25)
sec.footer_distance = Inches(.28)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Header', 'Footer']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.04
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
title = doc.styles['Title']
title.font.size = Pt(20)
title.font.bold = True
title.paragraph_format.space_after = Pt(6)
for name in ['Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.size = Pt(12)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(9)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.keep_with_next = True

def para(text, boldlead=None, size=None):
    p = doc.add_paragraph()
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold = True
        p.add_run(text[len(boldlead):])
    else:
        p.add_run(text)
    if size:
        for r in p.runs: r.font.size = Pt(size)
    return p

def heading(text): doc.add_heading(text, level=1)
def hyperlink(p, text, url):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '000000'); pr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); pr.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '18'); pr.append(sz)
    r.append(pr)
    t = OxmlElement('w:t'); t.text = text; r.append(t); h.append(r); p._p.append(h)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('MatterSyn  |  Research II draft  |  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
for r in footer.runs: r.font.size = Pt(9)

doc.add_paragraph('AI for Colloidal Materials Synthesis\nThe MatterSyn Research II Proposal', 'Title')
para('Principal investigator ____________________    CNetID ____________________', size=10)
para('University unit ____________________    Proposed allocation period ____________________', size=10)
para('Draft dated 20 September 2026. Provisional request: 375,000 service units and 1,000 GB persistent storage; CPU and GPU resources. Confirm billing assumptions before submission.', size=10)

heading('Research goals and significance')
para('Nanocrystals have technological importance spanning quantum dots in commercial displays [4] to nanodiamonds being developed for quantum sensing and computing [5]. Synthesizing nanocrystals with prescribed size, morphology, and defect characteristics remains challenging, including in strongly covalent materials such as diamond. Nanodiamonds with controlled color centers are particularly promising because these defects provide optical and spin properties useful for quantum technologies [5]. By proposing experimentally achievable synthesis recipes, MatterSyn aims to expand accessible nanomaterials and support advances in semiconductor, optoelectronic, and quantum technologies.')
para('Our initial focus is colloidal semiconductor nanocrystals, with extension to other material families as suitable evidence becomes available. We will propose recipes for specified composition, crystal phase, particle size, morphology, and properties. The central research question is whether explicit precursor chemistry, ordered processing steps, and sample-specific characterization improve recipe proposals over literature retrieval and text-only generation. MatterSyn connects a structured dataset to an interactive website so that researchers and computational models can trace each recipe, characterization result, and stated limitation to its source paper and supporting information.')
para('We will test three connected aims: construct a reproducible corpus of source-linked synthesis records; train compact models that predict precursors and processing conditions while representing uncertainty; and evaluate their proposals against held-out literature and, where feasible, experiments in a collaborating laboratory. Expected outputs are an auditable dataset, reproducible evaluation splits, model adapters, and a public materials atlas. Experimental validation is proposed work, not an achieved result.')

heading('Existing foundation and corpus scope')
para('MatterSyn version 0.33.0 contains 675 structured records, including 123 synthesis routes or variants, from 41 source contributions. The atlas covers 39 direct synthesis systems and 11 component collections. Records also include supporting procedures, observations, and benchmark rows; they are not 675 independent synthesis experiments. Thirty-six formal source readers document their supplied-source coverage. These results establish a working schema and publishing workflow [1].')
para('The fixed collection contains 9,532 provisional review groups, with three additional nested identity candidates held for resolution. A group is a candidate paper and associated files, not a verified recipe. Later arrivals will remain outside this campaign. Existing indexing and screening will be reused, but complete reading, extraction, and audit remain unfinished for most groups. Recipe yield and the number of training-eligible examples are not yet known.')

heading('Previous RCC use')
para('This project has not used a previous RCC allocation. There are therefore no project results or publications arising from prior RCC allocations to report. The existing MatterSyn pilot was developed outside RCC; no RCC performance or scaling measurements have yet been collected.')

doc.add_page_break()
heading('Computational research plan')
para('Corpus preparation and screening. We will deduplicate files using content hashes, assign stable source identifiers, and verify main-paper and supporting-information relationships using document content. Cached text and metadata will be reused. CPU jobs will perform parsing, selective OCR, layout analysis, and indexing. A lightweight screening stage will identify usable synthesis evidence and retain an explicit exclusion reason for papers without relevant recipes. Synthesis-rich sources with sample-resolved structural evidence will receive priority.', 'Corpus preparation and screening.')
para('Extraction and independent audit. Locally hosted language and vision models will extract chemicals, amounts, concentrations, reaction operations, atmosphere, temperature, duration, workup, storage, and characterization into a common schema. Each value will retain its units, source location, and sample assignment. An independent audit pass will revisit the source rather than merely approve the extractor’s output. Separate papers can run concurrently; an audit follows extraction for its own paper. Disagreements, missing information, and ambiguous figure assignments will enter a human review queue. Model agreement alone will not establish correctness.', 'Extraction and independent audit.')
para('Recipe learning. We will compare retrieval of similar recipes, structured prediction baselines, and parameter-efficient adaptation of approximately 1–8 billion parameter open-weight models. Inputs will initially be composition, reported phase, morphology, size, and available property targets. Outputs will be ranked precursor and operation sequences, supporting sources, and explicit uncertainty. The current collection has zero verified pairs linking measured atomic coordinates to a complete synthesis recipe. Exact structure-conditioned training will begin only if enough defensible pairs are recovered; reference unit cells will never be relabeled as measured product structures.', 'Recipe learning.')

heading('Evaluation and scientific verification')
para('A manually checked benchmark of 200 diverse paper groups will measure chemical identity, numerical values and units, operation order, missingness, and figure-to-sample assignment. A separate probability sample will estimate recipe prevalence and expected review effort; a deliberately diverse benchmark cannot estimate corpus yield. We will report errors and exclusions as well as retained records. Failed syntheses will be represented only when reported; the absence of a reported failure is not evidence of success.')
para('All records from a paper, its SI, and closely related recipe lineages will remain in the same training or evaluation split. Temporal and material-family holdouts, where data permit, will test generalization and limit leakage. We will compare precursor selection, condition prediction, constraint violations, citation support, and uncertainty calibration. Chemists will assess complete proposed procedures before any laboratory testing. Each accepted contribution will generate its website from the reviewed data, with checks on molecular depictions, apparatus conditions, and sample-specific characterization.')

heading('Efficient execution and staged milestones')
para('Weeks 1–2 will benchmark representative short, long, scanned, and SI-heavy bundles, measuring GPU memory, throughput, CPU utilization, and charged SUs. We will use Slurm job arrays, bounded concurrency, content-addressed caches, length-batched inference, and resumable checkpoints. Most GPU tasks require one device; small two-GPU experiments will proceed only if measured speedup justifies them. Job requests will be adjusted to observed memory and utilization rather than reserving entire nodes unnecessarily.')
para('Weeks 3–8 will prioritize extraction and audit of synthesis-rich papers and publish completed contributions in batches. Eight weeks is an initial campaign target, contingent on measured throughput and human audit capacity, not a guarantee of exhaustive curation. Months 3–6 will support model comparisons and held-out evaluation; months 7–12 will support error analysis, reproducible releases, and laboratory follow-up if feasible. New papers will remain in a separate queue until this campaign is assessed.')

doc.add_page_break()
heading('Compute request and workload estimates')
para('We propose 375,000 SUs as a planning envelope. The workload below assumes up to 10,000 paper groups for preparation and screening, rounded from the fixed collection. The 4,000-group detailed-audit scenario is a capacity assumption, not an observed 40% recipe yield. All timings are estimates to be replaced by pilot measurements.')
rows = [
    ['Workload', 'Planning calculation', 'Raw resource demand'],
    ['Parsing and selective OCR', '10,000 groups × 0.5 h × 4 CPU cores', '20,000 CPU core h'],
    ['Indexing and evaluation', '500 jobs × 1 h × 8 CPU cores', '4,000 CPU core h'],
    ['Model screening and extraction', '10,000 groups × 0.15 GPU h', '1,500 GPU h'],
    ['Independent source audit', '4,000 groups × 0.25 GPU h', '1,000 GPU h'],
    ['Model adaptation and evaluation', '20 runs × 25 GPU h per run', '500 GPU h'],
    ['Total', 'CPU-only and GPU jobs counted separately', '24,000 CPU core h\n3,000 GPU h'],
]
table = doc.add_table(rows=0, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [2.05, 2.88, 2.07]
for c,w in zip(table.columns,widths): c.width = Inches(w)
tblpr = table._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for side in ['top','left','bottom','right','insideH','insideV']:
    b=OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
tblpr.append(borders)
for idx,vals in enumerate(rows):
    cells=table.add_row().cells
    trpr=table.rows[-1]._tr.get_or_add_trPr()
    no_split=OxmlElement('w:cantSplit'); trpr.append(no_split)
    if idx==0:
        repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
    for j,(c,v) in enumerate(zip(cells,vals)):
        c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr=c._tc.get_or_add_tcPr()
        margins=OxmlElement('w:tcMar')
        for side,val in [('top',70),('bottom',70),('left',85),('right',85)]:
            e=OxmlElement('w:'+side); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); margins.append(e)
        tcpr.append(margins)
        if idx==0:
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8E8E8'); tcpr.append(shade)
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
        if j==2: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(v); r.font.size=Pt(10); r.bold=(idx==0 or idx==6)
budget = para('Budget conversion. For planning only, we allow 1 SU per CPU core-hour and 100 SUs per GPU-hour, the latter including proportional host CPU and RAM. These are provisional allowances, not verified Midway3 rates. The calculation is (24,000 × 1 + 3,000 × 100) × 1.15 = 372,600 SUs, rounded to 375,000. The 15% reserve covers reruns and difficult documents. RCC’s published Midway2 formula explicitly does not establish Midway3 billing [3]. We will confirm rates and revise the request or workload before submission.', 'Budget conversion.')
budget.paragraph_format.space_before = Pt(7)
para('Resource configuration. CPU jobs will use 4–8 cores and 16–32 GB RAM on one node. GPU inference jobs will use one GPU, 4–8 host cores, and 32–64 GB host RAM; adaptation may use two GPUs on one node and 64–128 GB host RAM. We target 24–48 GB device memory, using smaller or quantized models if needed. Every device counts toward GPU-hours. Initial concurrency will be capped at four CPU jobs and four single-GPU jobs, subject to RCC scheduling. No 768 GB big-memory nodes or DFT calculations are requested.', 'Resource configuration.')

heading('Storage and reproducible outputs')
para('We request 1,000 GB persistent storage: 150 GB for permitted source documents and metadata, 150 GB for structured records and selected evidence, 250 GB for model weights and environments, 350 GB for checkpoints and evaluation outputs, and 100 GB reserve. These are capacity estimates. Temporary page renders and intermediate files will use approximately 1 TB peak scratch space, subject to RCC policy, with deletion after validation. Deduplication, checkpoint retention limits, and selective rendering will control growth.')
para('Public releases will include permitted structured data, code, citations, audit history, and model artifacts where licenses allow. Original articles and SI will remain access-controlled and will not be placed in public repositories. We will respect source access and model licenses and retain versioned provenance for each released record.')

heading('References and project resources')
p=para('[1] MatterSyn dataset 0.33.0 and project status, 20 September 2026. ',size=9)
hyperlink(p,'Materials atlas','https://cuiyist.github.io/mattersyn-site/')
p.add_run(' and '); hyperlink(p,'code and audit history','https://github.com/cuiyist/mattersyn')
p=para('[2] University of Chicago RCC. ',size=9)
hyperlink(p,'Research II allocation requirements','https://rcc.uchicago.edu/accounts-allocations/research-ii-allocation-request')
p=para('[3] University of Chicago RCC. ',size=9)
hyperlink(p,'Calculations of service units','https://rcc.uchicago.edu/accounts-allocations/calculations-service-units')
p=para('[4] Nanosys. ',size=9)
hyperlink(p,'Quantum Dots Take Center Stage at Display Week 2025','https://www.nanosys.com/blog-newsroom/quantum-dots-take-center-stage-at-display-week-2025')
p.add_run('. 30 May 2025.')
p=para('[5] Liang J et al. ',size=9)
hyperlink(p,'Bottom-up synthesis of molecular nanodiamond from nanographene','https://doi.org/10.1038/s41586-026-10669-3')
p.add_run('. Nature 655, 102–108 (2026).')
for p in doc.paragraphs:
    if p.text.startswith(('[1]','[2]','[3]','[4]','[5]')):
        p.paragraph_format.space_after=Pt(2)
        for r in p.runs: r.font.size=Pt(9)

doc.core_properties.title='AI for Colloidal Materials Synthesis — MatterSyn Research II Proposal'
doc.core_properties.subject='Draft RCC allocation proposal'
doc.core_properties.author='MatterSyn project'
doc.core_properties.keywords='MatterSyn, synthesis, colloidal nanocrystals, RCC, Research II'
doc.save(OUT)
print(str(OUT))
print(f'File size: {OUT.stat().st_size:,} bytes')
print(f'Paragraph and table words: {sum(len(p.text.split()) for p in doc.paragraphs)+sum(len(c.text.split()) for row in table.rows for c in row.cells)}')
