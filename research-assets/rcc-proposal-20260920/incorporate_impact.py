from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

root = Path(__file__).parent
path = Path(r'[local path redacted]')
backup = root / 'MatterSyn_RCC_Research_II_Proposal_before_impact.docx'
if not backup.exists():
    copy2(path, backup)
doc = Document(path)
replacements = [
    ('We seek to develop an evidence-based system', 'Nanocrystals have technological importance spanning quantum dots in commercial displays [4] to nanodiamonds being developed for quantum sensing and computing [5]. Synthesizing nanocrystals with prescribed size, morphology, and defect characteristics remains challenging, including in strongly covalent materials such as diamond. Nanodiamonds with controlled color centers are particularly promising because these defects provide optical and spin properties useful for quantum technologies [5]. By proposing experimentally achievable synthesis recipes, MatterSyn aims to expand accessible nanomaterials and support advances in semiconductor, optoelectronic, and quantum technologies.'),
    ('The central research question is whether', 'Our initial focus is colloidal semiconductor nanocrystals, with extension to other material families as suitable evidence becomes available. We will propose recipes for specified composition, crystal phase, particle size, morphology, and properties. The central research question is whether explicit precursor chemistry, ordered processing steps, and sample-specific characterization improve recipe proposals over literature retrieval and text-only generation. MatterSyn connects a structured dataset to an interactive website so that researchers and computational models can trace each recipe, characterization result, and stated limitation to its source paper and supporting information.'),
]
for prefix, new_text in replacements:
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(matches) == 1, prefix
    p = matches[0]
    p.clear()
    p.add_run(new_text)

def link(p, label, url):
    h=OxmlElement('w:hyperlink')
    h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    for tag,value in [('color','000000'),('u','single'),('sz','18')]:
        e=OxmlElement('w:'+tag); e.set(qn('w:val'),value); pr.append(e)
    r.append(pr); t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); p._p.append(h)

refs = [
    ('[4] Nanosys. ', 'Quantum Dots Take Center Stage at Display Week 2025', 'https://www.nanosys.com/blog-newsroom/quantum-dots-take-center-stage-at-display-week-2025', '. 30 May 2025.'),
    ('[5] Liang J et al. ', 'Bottom-up synthesis of molecular nanodiamond from nanographene', 'https://doi.org/10.1038/s41586-026-10669-3', '. Nature 655, 102–108 (2026).'),
]
for lead,label,url,tail in refs:
    assert not any(p.text.startswith(lead) for p in doc.paragraphs)
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(2)
    p.add_run(lead).font.size=Pt(9)
    link(p,label,url)
    p.add_run(tail).font.size=Pt(9)
doc.save(path)

# Keep the original reproducible builder aligned with this local revision.
builder = root / 'build_proposal.py'
source = builder.read_text(encoding='utf-8')
for prefix,new_text in replacements:
    lines=source.splitlines()
    candidates=[i for i,line in enumerate(lines) if line.startswith("para('"+prefix)]
    assert len(candidates)==1, prefix
    lines[candidates[0]]='para('+repr(new_text)+')'
    source='\n'.join(lines)+'\n'
ref_code = """p=para('[4] Nanosys. ',size=9)
hyperlink(p,'Quantum Dots Take Center Stage at Display Week 2025','https://www.nanosys.com/blog-newsroom/quantum-dots-take-center-stage-at-display-week-2025')
p.add_run('. 30 May 2025.')
p=para('[5] Liang J et al. ',size=9)
hyperlink(p,'Bottom-up synthesis of molecular nanodiamond from nanographene','https://doi.org/10.1038/s41586-026-10669-3')
p.add_run('. Nature 655, 102–108 (2026).')
"""
source=source.replace("for p in doc.paragraphs:\n    if p.text.startswith(('[1]','[2]','[3]')):", ref_code+"for p in doc.paragraphs:\n    if p.text.startswith(('[1]','[2]','[3]','[4]','[5]')):")
builder.write_text(source,encoding='utf-8')
print('Updated',path)
print('Preserved original at',backup)
print('File size',path.stat().st_size)
